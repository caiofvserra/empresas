#!/usr/bin/env python3
"""Import the supplied Confea workbook as a rules/evidence catalog, never live data.

Requires openpyxl. Formulae are preserved, not evaluated by this importer. Values
previously filled in the workbook are isolated under workbookEvidence and cannot
be used as a current CREA-SP indicator value.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, time
import hashlib
import json
from pathlib import Path
import re

import openpyxl


DEFAULT_SOURCE = Path(
    "/tmp/codex-remote-attachments/01a122d1-7ab4-72d9-a20c-d6fc0fb6b033/"
    "4371acd3-2dbe-4c91-a185-ac1328a8df58/"
    "3-PL-2260-2023_PL-0926-2025_ANUAL_Relação-Dados-e-Indicadores-Resultados-Sistema-AJUSTE-INDICADOR-53-2-.xlsx"
)
DISPLAY_NAME = (
    "PL 2260 2023_PL 0926 2025_ANUAL_Relação Dados e Indicadores Resultados Sistema - AJUSTE INDICADOR 53 (2).xlsx"
)
REF_RE = re.compile(r"'([^']+)'!\$?([A-Z]+)\$?(\d+)")


def json_value(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return value


def color_evidence(color):
    if color is None:
        return None
    return {"type": color.type, "value": color.index, "tint": color.tint}


def code_of(value, previous: int | None = None) -> str:
    """Resolve only the workbook's simple sequential ID expressions; no Excel eval."""
    text = str(value)
    if text.startswith("="):
        if not re.fullmatch(r"=A\d+\+1", text) or previous is None:
            raise ValueError(f"Unsupported identifier expression {text!r}")
        return str(previous + 1)
    match = re.match(r"^(\d+(?:-[A-Z])?)", text)
    if not match:
        raise ValueError(f"Invalid metric/indicator identifier {text!r}")
    return match.group(1)


def classify_metric(sheet: str, code: str, program: str, theme: str, label: str) -> dict:
    text = f"{program} {theme} {label}".lower()
    if sheet == "Dados - Confea":
        return {"domain": "national_auxiliary", "use": "confea_reporting_dependency",
                "reason": "Competência nacional/terceira instância; não representa a população paulista."}
    if code == "37-C":
        return {"domain": "retired", "use": "audit_only",
                "reason": "A própria planilha identifica a métrica como excluída."}
    if "julgamento" in program.lower():
        return {"domain": "judgment", "use": "accountability_only",
                "reason": "Prestação de contas de processos/decisões; não é sinal de fiscalização digital por cruzamento."}
    if "fiscal" in program.lower():
        if any(word in text for word in ("conab", "aneel", "esgotamento", "municípios")):
            return {"domain": "digital_inspection", "use": "external_universe_coverage",
                    "reason": "Pode reconciliar o universo territorial/setorial externo com a cobertura documentada; não presume infração."}
        return {"domain": "inspection_accountability", "use": "accountability_only",
                "reason": "Resultado/recurso de fiscalização exigido no reporte; separado da descoberta de indícios digitais."}
    if "registro" in program.lower():
        if any(word in text for word in ("arts", "cat", "acervo")):
            domain = "art_and_acervo"
        elif any(word in text for word in ("empresa", "qt")):
            domain = "business"
        else:
            domain = "professionals"
        return {"domain": domain, "use": "decision_support_and_reporting",
                "reason": "Universo/atividade cadastral com recorte e unidade explícitos; preservar distinções registro, visto e pessoa distinta."}
    if any(word in text for word in ("receita", "despesa", "anuidade", "multa", "orçamento", "salários", "valor", "verbas")):
        return {"domain": "finance", "use": "financial_governance",
                "reason": "Fonte financeira homologada e recorte de exercício; sem usar contagem de débitos como contagem de pessoas."}
    return {"domain": "governance_auxiliary", "use": "accountability_context",
            "reason": "Informação de suporte/liderança útil ao reporte; não ocupa a visão principal sem pergunta decisória."}


def classify_indicator(identifier: int, scope: str) -> dict:
    if scope == "confea_national":
        return {"domain": "national_auxiliary", "use": "national_dependency_only",
                "reason": "Indicador exclusivamente Confea; não apresentar como medida do CREA-SP."}
    if scope == "mixed_requires_confea":
        return {"domain": "judgment", "use": "reporting_with_national_dependency",
                "reason": "Fórmula soma instâncias CREA e Confea; não calcular só com a parcela paulista nem chamar de total SP."}
    if identifier <= 2:
        domain, purpose = "registry", "Evolução do universo ativo comparável entre encerramentos de exercício."
    elif identifier == 3:
        domain, purpose = "technical_links", "Dimensionamento do QT/RT; a média não substitui cobertura de empresas com vínculo vigente."
    elif identifier <= 7:
        domain, purpose = "art_and_acervo", "Atividade formal documentada/acervo; não comprova emprego, renda ou infração individual."
    elif identifier <= 10 or identifier == 35:
        domain, purpose = "finance", "Regularidade financeira e cobrança com denominador de registro ativo/exercício explícito."
    elif identifier <= 16:
        domain, purpose = "inspection_accountability", "Indicador de reporte institucional; manter na prestação de contas, separado dos sinais digitais."
    elif identifier <= 48:
        domain, purpose = "judgment", "Indicador normativo de processos/decisões; manter na prestação de contas sem confundir com triagem digital."
    elif identifier == 49:
        domain, purpose = "registration_service", "Prazo de registro PF/PJ/CAT/CAO conforme a fórmula anual, não SLA inventado por protocolo."
    elif identifier <= 52:
        domain, purpose = "digital_inspection", "Cobertura de universo externo definido; permite localizar lacunas para verificação sem inferir irregularidade."
    else:
        domain, purpose = "inspection_budget", "Participação do orçamento aprovado destinada à fiscalização; não prova aumento real acima da inflação."
    return {"domain": domain, "use": "decision_support_and_reporting" if identifier <= 10 or identifier >= 49 else "accountability_only",
            "reason": purpose}


def run(source: Path, destination: Path, ui_destination: Path | None = None) -> dict:
    workbook = openpyxl.load_workbook(source, data_only=False)
    cached = openpyxl.load_workbook(source, data_only=True)
    evidence_sheets = []
    for sheet in workbook:
        cells = []
        for cell in sorted(sheet._cells.values(), key=lambda c: (c.row, c.column)):
            if cell.value is None:
                continue
            item = {
                "cell": cell.coordinate,
                "value": json_value(cell.value),
                "valueType": cell.data_type,
                "numberFormat": cell.number_format,
                "fill": {"pattern": cell.fill.patternType, "foreground": color_evidence(cell.fill.fgColor)},
                "strikethrough": bool(cell.font.strike),
            }
            if cell.data_type == "f":
                item["formula"] = cell.value
                item["cachedValue"] = json_value(cached[sheet.title][cell.coordinate].value)
                item["cachedValueIsLiveData"] = False
            if cell.comment:
                item["comment"] = {"text": cell.comment.text, "author": cell.comment.author}
            if cell.hyperlink:
                item["hyperlink"] = cell.hyperlink.target
            cells.append(item)
        evidence_sheets.append({
            "name": sheet.title, "maxRow": sheet.max_row, "maxColumn": sheet.max_column,
            "state": sheet.sheet_state, "mergedRanges": [str(r) for r in sheet.merged_cells.ranges],
            "protected": bool(sheet.protection.sheet), "autoFilterRange": sheet.auto_filter.ref,
            "freezePanes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
            "hiddenRows": [n for n, d in sheet.row_dimensions.items() if d.hidden],
            "hiddenColumns": [n for n, d in sheet.column_dimensions.items() if d.hidden],
            "cells": cells,
        })

    metrics = []
    metric_by_ref = {}
    for sheet_name, prefix, owner in (("Dados - Crea", "crea", "Crea"), ("Dados - Confea", "confea", "Confea")):
        sheet = workbook[sheet_name]
        for row in range(5, sheet.max_row + 1):
            if sheet.cell(row, 1).value is None or sheet.cell(row, 5).value is None:
                continue
            val = lambda column: json_value(sheet.cell(row, column).value)
            code = code_of(val(1))
            related = val(12)
            indicators_related = [int(v) for v in re.findall(r"\d+", str(related))] if related is not None else []
            deleted = code == "37-C"
            fill = sheet.cell(row, 1).fill.fgColor
            highlight = "yellow" if fill.type == "theme" and fill.theme == 7 else "green" if fill.type == "theme" and fill.theme == 9 else "unmarked"
            metric = {
                "id": f"{prefix}:{code}", "code": code, "codeOriginal": val(1),
                "label": val(5), "subprogram": val(2), "theme": val(3), "basis": val(4),
                "temporalCut": val(6), "definition": val(7), "orientation": val(8),
                "unit": val(9), "valueFormat": val(9), "countingUnitRequiresDefinitionReview": True,
                "verificationFormula": val(11),
                "relatedIndicatorIds": indicators_related, "relatedIndicatorIdsOriginal": related,
                "reportPresentation": val(13), "tcuDataCategory": val(14), "justification": val(15),
                "additionalNotes": [{"cell": sheet.cell(row, c).coordinate, "text": val(c)}
                                    for c in range(16, sheet.max_column + 1) if val(c) is not None],
                "owner": owner, "scope": "crea_sp" if owner == "Crea" else "confea_national",
                "catalogStatus": "retired" if deleted else "active_rule",
                "workbookHighlight": highlight,
                "reporting": {"annualConsolidation": True,
                              "monthlyCollectionIndicated": owner == "Crea" and highlight in ("yellow", "green") and not deleted,
                              "monthlyRuleEvidence": "INSTRUÇÃO!C7" if owner == "Crea" and highlight in ("yellow", "green") else None,
                              "firstSubmission": "2026-04-01" if "(novo)" in str(val(1)) and owner == "Crea" else None},
                "sourceConnectionRequired": True, "status": "retired" if deleted else "source_connection_required",
                "numericValue": None,
                "sourceRoute": {
                    "transport": "on_premises_gateway" if owner == "Crea" else "official_database_or_api_required",
                    "databases": ["SQL Server", "DB2"] if owner == "Crea" else [],
                    "physicalMappingStatus": "requires_schema_homologation",
                    "workbookIsDataFeed": False,
                },
                "relevance": classify_metric(sheet_name, code, str(val(2)), str(val(3)), str(val(5))),
                "workbookEvidence": {"sheet": sheet_name, "row": row, "idCell": f"A{row}",
                                     "labelCell": f"E{row}", "valueCell": f"J{row}",
                                     "filledValueIsLiveData": False},
            }
            metrics.append(metric)
            metric_by_ref[(sheet_name, f"J{row}")] = metric

    indicators = []
    sheet = workbook["Indicadores RG"]
    previous = None
    for row in range(5, sheet.max_row + 1):
        if sheet.cell(row, 1).value is None or sheet.cell(row, 3).value is None:
            continue
        val = lambda column: json_value(sheet.cell(row, column).value)
        identifier = int(code_of(val(1), previous))
        previous = identifier
        formula = val(7)
        dependencies = []
        for ref_sheet, column, ref_row in REF_RE.findall(str(formula)):
            cell = f"{column}{ref_row}"
            metric = metric_by_ref.get((ref_sheet, cell))
            dependency = {"sheet": ref_sheet, "cell": cell, "metricKey": metric["id"] if metric else None,
                          "metricCode": metric["code"] if metric else None,
                          "metricLabel": metric["label"] if metric else None,
                          "temporalCut": metric["temporalCut"] if metric else None}
            if dependency not in dependencies:
                dependencies.append(dependency)
        owner = val(8)
        scope = "mixed_requires_confea" if "\n" in str(owner) else "confea_national" if owner == "Confea" else "crea_sp"
        indicators.append({
            "id": f"confea-indicator:{identifier}", "indicatorId": identifier, "codeOriginal": val(1),
            "label": val(3), "subprogram": val(2), "objective": val(4), "formulaText": val(5),
            "excelFormula": formula, "unit": val(6), "owner": owner, "scope": scope,
            "reportPresentation": val(9), "systemResultsPresentation": val(10),
            "tcuDemand": val(11), "notes": val(12), "amendmentNote": val(13),
            "dependencies": dependencies, "sourceConnectionRequired": True,
            "status": "source_connection_required", "numericValue": None,
            "publicationAssessment": {
                "workbookMarksReportMandatory": val(9) == "Obrigatória",
                "legalPublicationIndependentlyVerified": False,
                "scope": "aggregated_report_only_pending_legal_validation",
                "containsPermissionToPublishPersonalData": False,
            },
            "relevance": classify_indicator(identifier, scope),
            "workbookEvidence": {"sheet": sheet.title, "row": row, "idCell": f"A{row}",
                                 "labelCell": f"C{row}", "formulaCell": f"G{row}",
                                 "filledValueIsLiveData": False},
        })

    instructions = [{"cell": c.coordinate, "text": c.value} for c in sorted(workbook["INSTRUÇÃO"]._cells.values(), key=lambda c: (c.row, c.column))
                    if c.value is not None]
    document = {
        "schemaVersion": "1.0.0", "catalogVersion": "v8", "kind": "indicator_rules_catalog",
        "source": {"displayName": DISPLAY_NAME, "fileName": source.name,
                   "sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "bytes": source.stat().st_size,
                   "providedBy": "user", "sourceType": "xlsx_rules_and_historical_evidence",
                   "jurisdictionForActiveDashboard": "CREA-SP / São Paulo", "isAutomaticDataSource": False},
        "refreshContract": {
            "numericValuesFromWorkbookEnabled": False,
            "allowedLiveSources": ["SQL Server via gateway", "DB2 via gateway", "approved official API"],
            "annualReportingDate": {"day": 1, "month": 4, "evidence": "INSTRUÇÃO!C4/C6"},
            "monthlyConsolidation": {"appliesTo": "yellow_and_green_metrics", "evidence": "INSTRUÇÃO!C7"},
            "initialNewMetricsSubmissionDate": "2026-04-01",
            "missingValuePolicy": "null_is_unknown_not_zero", "evidence": "INSTRUÇÃO!C17",
            "computedIndicatorsOnlyAfterSourceAndFormulaHomologation": True,
        },
        "summary": {
            "sheets": len(evidence_sheets), "creaMetrics": sum(m["owner"] == "Crea" for m in metrics),
            "confeaMetrics": sum(m["owner"] == "Confea" for m in metrics), "indicators": len(indicators),
            "creaMetricsMarkedReportMandatory": sum(m["owner"] == "Crea" and m["reportPresentation"] == "Obrigatória" for m in metrics),
            "creaMetricsWithTcuCategory": sum(m["owner"] == "Crea" and m["tcuDataCategory"] is not None for m in metrics),
            "indicatorScopes": dict(Counter(i["scope"] for i in indicators)),
            "liveNumericValues": 0,
        },
        "instructions": instructions, "metrics": metrics, "indicators": indicators,
        "auditNotes": [
            {"id": "indicator_53_amendment", "severity": "interpretation_guard", "evidence": ["Indicadores RG!G57", "Indicadores RG!M57", "Dados - Crea!E189", "Dados - Crea!E190"],
             "message": "Fórmula ajustada =J190/J189: código183 (orçamento aprovado destinado à fiscalização) / código182 (orçamento total aprovado). A nota chama 189/190 de dados; são números de linha, não códigos."},
            {"id": "indicator_53_not_real_growth", "severity": "requires_separate_measure", "evidence": ["Indicadores RG!E57", "Indicadores RG!L57"],
             "message": "Fórmula mede participação orçamentária no exercício. Não mede por si aumento10% acima da inflação até2027; isso exige orçamento de base, orçamento comparável e índice de preços homologado."},
            {"id": "indicator_35_denominator_mismatch", "severity": "requires_formula_homologation", "evidence": ["Indicadores RG!E39", "Indicadores RG!G39"],
             "message": "Texto usa multas não pagas no exercício anterior como denominador; Excel usa J156 (exercício atual), enquanto o anterior é J155. Ambas versões preservadas; nenhuma correção silenciosa."},
            {"id": "new_metrics_instruction_range", "severity": "source_inconsistency", "evidence": ["INSTRUÇÃO!C6", "Dados - Crea!A190"],
             "message": "Instrução menciona novas métricas162a182; folha Dados - Crea inclui também183. Catálogo conserva183."},
            {"id": "metric_37c_retired", "severity": "retired_rule", "evidence": ["Dados - Crea!A44", "Dados - Crea!P44"],
             "message": "Métrica37-C está marcada excluída. Conservada para auditoria, sem valor ou uso operacional."},
            {"id": "indicator_44_wrong_observation", "severity": "source_inconsistency", "evidence": ["Indicadores RG!G48", "Indicadores RG!L48"],
             "message": "Fórmula do indicador44 aponta Dados - Confea!J71; observação menciona Dados-Crea. Preservar a dependência efetiva e homologar redação."},
            {"id": "report_obligation_is_not_raw_publication", "severity": "publication_guard", "evidence": ["Indicadores RG!I4", "Indicadores RG!J4", "Dados - Crea!M4"],
             "message": "Marca Obrigatória é evidência da exigência no Relatório de Gestão/caderno indicada na planilha. Não autoriza publicar microdados, débitos individuais, nomes ou acessos internos."},
            {"id": "workbook_values_not_live", "severity": "data_source_guard", "evidence": ["Dados - Crea!J4", "INSTRUÇÃO!C7"],
             "message": "Números preenchidos e resultados em cache são evidência histórica do arquivo recebido, sem competência validada nesta importação. Nunca funcionam como feed atualizado ou medida ativa."},
        ],
        "workbookEvidence": {"policy": "read_only_source_evidence_not_current_dashboard_values",
                             "themeXml": workbook.loaded_theme.decode("utf-8"), "sheets": evidence_sheets},
    }
    assert len(metrics) == 281 and len(indicators) == 53
    assert len({m["id"] for m in metrics}) == len(metrics)
    assert [i["indicatorId"] for i in indicators] == list(range(1, 54))
    assert all(d["metricKey"] for i in indicators for d in i["dependencies"])
    assert all(i["numericValue"] is None for i in indicators)
    assert indicators[-1]["excelFormula"] == "='Dados - Crea'!J190/'Dados - Crea'!J189"
    assert [d["metricCode"] for d in indicators[-1]["dependencies"]] == ["183", "182"]
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if ui_destination:
        def public_rule_projection(item):
            # This block contains ONLY cell addresses, not their raw values.
            evidence = item["workbookEvidence"]
            return {**{k: v for k, v in item.items() if k != "workbookEvidence"}, "evidence": evidence}

        ui_projection = {
            "schemaVersion": document["schemaVersion"], "catalogVersion": document["catalogVersion"],
            "kind": "indicator_rules_ui_projection", "source": document["source"],
            "refreshContract": document["refreshContract"], "summary": document["summary"],
            "instructions": document["instructions"], "auditNotes": document["auditNotes"],
            "metrics": [public_rule_projection(m) for m in metrics if m["owner"] == "Crea"],
            "indicators": [public_rule_projection(i) for i in indicators],
        }
        assert len(ui_projection["metrics"]) == 186
        assert "workbookEvidence" not in json.dumps(ui_projection, ensure_ascii=False)
        ui_destination.parent.mkdir(parents=True, exist_ok=True)
        ui_destination.write_text(json.dumps(ui_projection, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
        summary_path = ui_destination.with_name("confea-summary-v8.json")
        summary_path.write_text(json.dumps(document["summary"], ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    return document["summary"]


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "src/domain/confea-catalog-v8.json")
    parser.add_argument("--ui-output", type=Path, default=Path(__file__).resolve().parents[1] / "src/domain/confea-ui-v8.json")
    args = parser.parse_args()
    print(json.dumps(run(args.source, args.output, args.ui_output), ensure_ascii=False))
